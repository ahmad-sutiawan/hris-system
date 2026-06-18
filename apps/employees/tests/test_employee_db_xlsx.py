from pathlib import Path

from django.conf import settings
from django.test import TestCase

from apps.core.models import Plant, Tenant
from apps.employees.models import Employee
from apps.employees.services.import_talenta import import_employees_talenta_xlsx
from apps.employees.services.talenta_validate import validate_talenta_xlsx
from apps.employees.talenta_vocabulary import TALENTA_COLUMNS


EMPLOYEE_DB = Path(settings.MEDIA_ROOT) / "employee_db.xlsx"


class EmployeeDbXlsxTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._has_file = EMPLOYEE_DB.is_file()

    def setUp(self):
        self.tenant = Tenant.objects.create(name="BPS", slug="bps-db-test")

    def test_employee_db_schema_fits_model(self):
        if not self._has_file:
            self.skipTest(f"{EMPLOYEE_DB} tidak ditemukan")

        content = EMPLOYEE_DB.read_bytes()
        validation = validate_talenta_xlsx(content)
        self.assertTrue(
            validation.ok,
            "\n".join(validation.errors[:20]),
        )
        self.assertEqual(validation.row_count, 695)

    def test_employee_db_import_and_relations(self):
        if not self._has_file:
            self.skipTest(f"{EMPLOYEE_DB} tidak ditemukan")

        content = EMPLOYEE_DB.read_bytes()
        result = import_employees_talenta_xlsx(self.tenant, content)
        self.assertEqual(result["errors"], [], result["errors"][:5])

        employees = Employee.objects.filter(tenant=self.tenant)
        self.assertEqual(employees.count(), 695)

        for emp in employees.select_related("plant", "department", "job_position", "job_level"):
            self.assertIsNotNone(emp.plant_id)
            self.assertIsNotNone(emp.department_id)
            self.assertIsNotNone(emp.job_position_id)
            self.assertIsNotNone(emp.job_level_id)
            self.assertEqual(emp.department.plant_id, emp.plant_id)
            self.assertEqual(emp.job_position.plant_id, emp.plant_id)
            emp.full_clean()

        self.assertGreaterEqual(employees.filter(manager__isnull=False).count(), 600)
        self.assertGreaterEqual(Plant.objects.filter(tenant=self.tenant).count(), 3)

    def test_employee_db_round_trip_core_fields(self):
        if not self._has_file:
            self.skipTest(f"{EMPLOYEE_DB} tidak ditemukan")

        import openpyxl

        content = EMPLOYEE_DB.read_bytes()
        import_employees_talenta_xlsx(self.tenant, content)

        wb = openpyxl.load_workbook(EMPLOYEE_DB, read_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

        mismatches = 0
        for row in ws.iter_rows(min_row=2, max_row=12, values_only=True):
            excel = dict(zip(headers, row))
            emp = Employee.objects.get(
                tenant=self.tenant,
                employee_id=str(excel["Employee ID"]).strip(),
            )
            exported = emp.talenta_export_row()
            for col in TALENTA_COLUMNS:
                ex = str(excel.get(col) or "").strip()
                out = str(exported.get(col, "") or "").strip()
                if col in ("Residential Address", "Citizen ID Address"):
                    ex = ex.replace("\n", " ")
                if ex and ex != out:
                    mismatches += 1
        self.assertEqual(mismatches, 0)
