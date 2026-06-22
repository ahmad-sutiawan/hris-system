"""Shared Excel (.xlsx) import/export helpers."""

from __future__ import annotations

import json
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any, Callable, Iterable

from django.http import HttpResponse

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_ALIASES = frozenset({"xlsx", "csv"})


def export_format_requested(request) -> bool:
    return request.GET.get("export") in EXPORT_ALIASES


def cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            from django.utils import timezone

            value = timezone.localtime(value)
        return value.replace(tzinfo=None).isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_xlsx(headers: list[str], rows: Iterable[list[Any]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append([cell_value(value) for value in row])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_xlsx_rows(file_bytes: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header_cells = next(iterator, None)
    if not header_cells:
        workbook.close()
        return [], []

    headers = [str(header).strip() if header is not None else "" for header in header_cells]
    records: list[dict[str, Any]] = []
    for cells in iterator:
        if cells is None:
            continue
        if all(cell is None or str(cell).strip() == "" for cell in cells):
            continue
        record: dict[str, Any] = {}
        for index, key in enumerate(headers):
            if not key:
                continue
            record[key] = cells[index] if index < len(cells) else None
        records.append(record)
    workbook.close()
    return headers, records


def xlsx_http_response(content: bytes, filename: str) -> HttpResponse:
    if not filename.lower().endswith(".xlsx"):
        stem = filename.rsplit(".", 1)[0]
        filename = f"{stem}.xlsx"
    response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def queryset_to_xlsx(qs, headers: list[str], row_builder: Callable) -> bytes:
    return write_xlsx(headers, (row_builder(row) for row in qs))


def normalize_xlsx_filename(filename: str) -> str:
    if filename.lower().endswith(".xlsx"):
        return filename
    return f"{filename.rsplit('.', 1)[0]}.xlsx"
